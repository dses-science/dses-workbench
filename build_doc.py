"""Build a DSES-styled PDF from a Markdown source (e.g. Installing.md).

The deliverable is the PDF:
    DSES_Radio_Astronomy_Workbench_Installation.pdf

Internally the document is rendered to a temporary .docx (python-docx, styled
to match the DSES house template — margins, title color, Heading 1
white-on-teal banner, Heading 3 dark-teal, page header with "DSES" + subtitle,
footer with date + page number) and then converted to PDF. The .docx is a
throwaway intermediate — it is written to a temp file and removed after the
PDF is built, unless you pass --docx to keep it. PDF conversion uses
LibreOffice on macOS/Linux and Microsoft Word on Windows (see
convert_docx_to_pdf). On Windows the PDF is printed through a PDF printer
driver, never Word's own exporter (it rasterizes the house fonts): Acrobat
Distiller first (named font subsets), with an automatic, bounded fallback
to 'Microsoft Print to PDF' when Acrobat's font-capture helper crashes or
the print stalls; DSES_PDF_ENGINE=msprint|distiller|pdfmaker forces one.
The build must never stop on a dialog: see _word_app / _DialogSentinel.

Handles the Markdown subset used in Installing.md:
    # / ## / ###       — headings
    paragraphs         — regular text with inline **bold**, *italic*, `code`
    - / *              — bulleted lists
    1.                 — numbered lists
    ```                — fenced code blocks (monospace, light-gray background)
    | a | b |          — tables (with --- alignment separator)
    [text](url)        — hyperlinks (rendered as text with the URL after, since
                         full hyperlink XML in python-docx is verbose)

Run from the project root with the project's conda env:
    .conda\\python.exe build_install_docx.py
"""

from __future__ import annotations

import html
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Smart-quote pre-pass: convert ASCII apostrophes and double quotes in body
# text to typographer equivalents, leaving fenced code blocks and inline
# `code` spans untouched. Matches DSES house style (Word's AutoCorrect would
# do the same thing for hand-typed text, but python-docx writes text
# verbatim and Word's autocorrect doesn't run on programmatic content).
# ---------------------------------------------------------------------------

LEFT_DOUBLE  = '“'   # "
RIGHT_DOUBLE = '”'   # "
APOSTROPHE   = '’'   # ’ (also doubles as right single quote)


def smartify_quotes(text: str) -> str:
    out = []
    in_fence = False
    dq_state = {'open': True}  # next " is opening
    for line in text.split('\n'):
        if line.lstrip().startswith('```'):
            in_fence = not in_fence
            out.append(line)
            continue
        if in_fence:
            out.append(line)
            continue
        # Reset double-quote state on paragraph breaks so an unmatched quote
        # in one paragraph doesn't flip the polarity for the next.
        if line.strip() == '':
            dq_state['open'] = True
        out.append(_smartify_line(line, dq_state))
    return '\n'.join(out)


def _smartify_line(line: str, dq_state: dict) -> str:
    result = []
    i = 0
    n = len(line)
    while i < n:
        c = line[i]
        if c == '`':
            # Inline code span — emit verbatim through the closing backtick.
            j = line.find('`', i + 1)
            if j == -1:
                result.append(c)
                i += 1
                continue
            result.append(line[i:j + 1])
            i = j + 1
            continue
        if c == "'":
            # Apostrophe only if it's after a letter (contraction / possessive).
            # Leave others straight — they might be opening single quotes,
            # feet markers, or other punctuation.
            prev = line[i - 1] if i > 0 else ''
            if prev.isalpha():
                result.append(APOSTROPHE)
            else:
                result.append(c)
            i += 1
            continue
        if c == '"':
            result.append(LEFT_DOUBLE if dq_state['open'] else RIGHT_DOUBLE)
            dq_state['open'] = not dq_state['open']
            i += 1
            continue
        result.append(c)
        i += 1
    return ''.join(result)


# Defaults for the install guide; CLI options can override. Used as module-
# level constants because the cover-page / header builders read them.
SRC = Path("Installing.md")
DST_PDF  = Path("DSES_Radio_Astronomy_Workbench_Installation.pdf")
DOC_TITLE    = "DSES Radio Astronomy Workbench"
DOC_SUBTITLE = "Installation Guide"
def _app_version_default():
    """'v<APP_VERSION>' read from dses_workbench.py next to this script, so a
    cover stamps the release being cut even when --version is not passed.
    (Before 27-Sep-2026 the default was a frozen "v1.1.6", and the release
    workflow PDF shipped with that stamp whenever the flag was forgotten.)"""
    try:
        text = (Path(__file__).resolve().parent / "dses_workbench.py").read_text(
            encoding="utf-8", errors="replace")
        m = re.search(r'^APP_VERSION\s*=\s*"([^"]+)"', text, re.M)
        if m:
            return "v" + m.group(1)
    except OSError:
        pass
    return "v1.1.6"


DOC_VERSION  = _app_version_default()
DOC_AUTHOR   = "Richard M Hambly (K0GD)"
DOC_ORG      = "DSES"

# Typography. Aligned with C:\CNS-Systems\DOCUMENT_STANDARDS.md section 3
# (Rick, 19-Aug-2026) so DSES documents match CNS Systems reports: Minion Pro
# body, Myriad Pro headings/display, Source Code Pro for code. The DSES skin
# (teal H1 banner, colors, sizes) is unchanged — only the faces moved.
#
# IMPORTANT: these are OpenType-PS (CFF) faces. Word's SaveAs-PDF silently
# RASTERIZES them (verified 19-Aug-2026: Minion/Myriad runs came out as images
# with no text layer, while TrueType Source Code Pro embedded fine), so the
# Windows converter below prints through a PDF printer driver instead
# (Adobe PDF/Distiller first, 'Microsoft Print to PDF' as the automatic
# fallback — see _pdf_engine). If you change these back to
# TrueType faces (Calibri/Cambria/Consolas), the plain SaveAs path is
# adequate again.
FONT_BODY = "Minion Pro"
FONT_HEAD = "Myriad Pro"
FONT_MONO = "Source Code Pro"

# DSES house style (pulled from EVE-26 + Pulsar installation reference docs)
TITLE_COLOR_RGB     = RGBColor(0x15, 0x60, 0x82)  # teal/blue
H1_BG_HEX           = "156082"                    # same teal as banner fill
H1_TEXT_RGB         = RGBColor(0xFF, 0xFF, 0xFF)  # white on the banner
H3_COLOR_RGB        = RGBColor(0x0A, 0x2F, 0x40)  # dark teal
SUBTITLE_COLOR_RGB  = RGBColor(0x59, 0x59, 0x59)  # mid-gray
CODE_FILL_HEX       = "F2F2F2"                    # very light gray
PAGE_MARGINS = dict(  # inches; matches both DSES reference documents
    left=1.00, right=0.81, top=1.36, bottom=1.00,
)
# Optional logo shown right-justified in the page header (pages 2+). Set via
# --header-logo; None keeps the classic text-only header.
HEADER_LOGO = None


# ---------------------------------------------------------------------------
# Inline run handling — split "...some **bold** and `code` text..." into a
# list of (text, style_dict) tuples.
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(
    r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))'
)
LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')


def parse_inline(text: str, base=None):
    """Yield (text, style) runs for one line of Markdown.

    Emphasis nests: a `code` span or a [link](url) inside **bold** or
    *italic* keeps the outer style. Before 4-Oct-2026 the inside of a bold or
    italic span was emitted verbatim, so **`launcher.bat`** printed its
    backticks in every guide (81 of them in the 1.6.0 install guide). The
    token patterns are unchanged, so nothing else renders differently; code
    spans stay literal (`**x**` prints its asterisks)."""
    base = dict(base or {})
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], dict(base)
        token = m.group(0)
        if token.startswith('**'):
            yield from parse_inline(token[2:-2], {**base, 'bold': True})
        elif token.startswith('`'):
            yield token[1:-1], {**base, 'code': True}
        elif token.startswith('['):
            lm = LINK_RE.match(token)
            if lm:
                label, url = lm.group(1), lm.group(2)
                if label.strip() == url.strip():
                    yield label, {**base, 'code': True}
                else:
                    yield label, dict(base)
                    yield f" ({url})", {**base, 'code': True}
        elif token.startswith('*'):
            yield from parse_inline(token[1:-1], {**base, 'italic': True})
        pos = m.end()
    if pos < len(text):
        yield text[pos:], dict(base)


def add_runs(paragraph, text: str):
    for chunk, style in parse_inline(text):
        if not chunk:
            continue
        chunk = html.unescape(chunk)
        run = paragraph.add_run(chunk)
        if style.get('bold'):
            run.bold = True
        if style.get('italic'):
            run.italic = True
        if style.get('code'):
            run.font.name = FONT_MONO
            run.font.size = Pt(9.5)


# ---------------------------------------------------------------------------
# DSES style primitives — small wrappers around the raw OOXML so the body
# of the document builder stays readable.
# ---------------------------------------------------------------------------

def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, color_hex):
    """Paint a solid fill behind a paragraph (used for the DSES H1 banner)."""
    pPr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)


def new_numbered_list(doc):
    """Return a fresh w:num id that restarts the 'List Number' sequence at 1.

    Word numbers every 'List Number' paragraph in one running sequence, so a
    second numbered list continues where the first stopped. A new w:num that
    points at the style's abstractNum with a startOverride of 1 restarts it.
    Returns None (caller keeps the plain style) if the lookup fails.
    """
    try:
        style_pPr = doc.styles['List Number'].element.pPr
        style_num = style_pPr.find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val'))
        numbering = doc.part.numbering_part.numbering_definitions._numbering
        abstract_id = None
        for num in numbering.findall(qn('w:num')):
            if num.get(qn('w:numId')) == style_num:
                abstract_id = num.find(qn('w:abstractNumId')).get(qn('w:val'))
                break
        if abstract_id is None:
            return None
        new_num = numbering.add_num(int(abstract_id))
        new_num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return new_num.numId
    except Exception:
        return None


def set_list_num(paragraph, num_id):
    """Point a list paragraph at the given w:num (level 0)."""
    if num_id is None:
        return
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:numPr')):
        pPr.remove(old)
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
    numId = OxmlElement('w:numId'); numId.set(qn('w:val'), str(num_id))
    numPr.append(ilvl); numPr.append(numId)
    pPr.append(numPr)


def add_page_number_field(paragraph):
    """Insert a PAGE field code so Word numbers pages live."""
    run = paragraph.add_run()
    for tag in ('begin', None, 'end'):
        if tag is None:
            instr = OxmlElement('w:instrText')
            instr.text = 'PAGE'
            run._r.append(instr)
        else:
            ch = OxmlElement('w:fldChar')
            ch.set(qn('w:fldCharType'), tag)
            run._r.append(ch)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Cover page + section header/footer (DSES house style)
# ---------------------------------------------------------------------------

def configure_section(section):
    section.left_margin   = Inches(PAGE_MARGINS['left'])
    section.right_margin  = Inches(PAGE_MARGINS['right'])
    section.top_margin    = Inches(PAGE_MARGINS['top'])
    section.bottom_margin = Inches(PAGE_MARGINS['bottom'])

    # Suppress the header/footer on the cover page: enable a distinct
    # first-page header/footer and leave it empty. The primary header/footer
    # configured below then applies only to page 2 onward.
    section.different_first_page_header_footer = True

    # Header (pages 2+): document title on line 1, subtitle on line 2.
    hdr = section.header
    # Reuse existing first paragraph; we control its content fully.
    p1 = hdr.paragraphs[0]
    p1.text = ''
    r = p1.add_run(DOC_TITLE)
    r.bold = True
    r.font.name = FONT_HEAD
    r.font.size = Pt(11)
    r.font.color.rgb = TITLE_COLOR_RGB
    # Optional logo, right-justified on the title line via a right tab stop
    # at the text-column edge.
    if HEADER_LOGO and Path(HEADER_LOGO).is_file():
        content_w = 8.5 - PAGE_MARGINS['left'] - PAGE_MARGINS['right']
        p1.paragraph_format.tab_stops.add_tab_stop(
            Inches(content_w), WD_TAB_ALIGNMENT.RIGHT)
        # The built-in Header style carries center (3.25") and right (6.5")
        # stops; a single tab would land on the center one and park the logo
        # mid-page. Emit w:val="clear" entries so only our stop remains.
        # NOTE: w:tab elements must be in ascending w:pos order or Word
        # discards the list — insert 9360 first, then 4680 in front of it.
        tabs_el = p1._p.pPr.find(qn('w:tabs'))
        for pos_twips in ('9360', '4680'):          # 6.5" then 3.25"
            clear = OxmlElement('w:tab')
            clear.set(qn('w:val'), 'clear')
            clear.set(qn('w:pos'), pos_twips)
            tabs_el.insert(0, clear)
        p1.add_run('\t')
        p1.add_run().add_picture(str(HEADER_LOGO), height=Inches(0.42))
    p2 = hdr.add_paragraph()
    r2 = p2.add_run(DOC_SUBTITLE)
    r2.italic = True
    r2.font.name = FONT_HEAD
    r2.font.size = Pt(9)
    r2.font.color.rgb = SUBTITLE_COLOR_RGB

    # Footer: date <tab> page number — matches the reference DSES docs.
    fp = section.footer.paragraphs[0]
    fp.text = ''
    today = date.today().strftime("%d-%b-%y")  # e.g. "16-May-26"
    fr = fp.add_run(today)
    fr.font.size = Pt(9)
    fr.font.color.rgb = SUBTITLE_COLOR_RGB
    fp.add_run('\t')
    add_page_number_field(fp)
    for run in fp.runs:
        run.font.name = FONT_HEAD
        run.font.size = Pt(9)
        run.font.color.rgb = SUBTITLE_COLOR_RGB


def _toc_plain(title):
    """Strip the small subset of inline markdown the headings use, for a clean
    table-of-contents entry."""
    t = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', title)   # [text](url) -> text
    return t.replace('**', '').replace('`', '').strip()


def _dedent(line, n):
    """Remove up to n leading spaces: the indent of the fence that opened a
    code block nested under a list item."""
    k = 0
    while k < n and k < len(line) and line[k] == ' ':
        k += 1
    return line[k:]


def check_fences(text, src_path):
    """Refuse a source with a damaged code fence.

    A fence needs three backticks. A line of exactly two (with or without a
    language word) is a fence that lost one, and everything between two such
    lines renders as run-on inline text with curly quotes and the language
    word glued on. Every code block of the install guide did, from the 1.1.8
    cut (3-Aug-2026) until this check was added (4-Oct-2026): nothing in the
    build complained."""
    bad = [n for n, ln in enumerate(text.split('\n'), 1)
           if re.match(r'^\s*``(?!`)\w*\s*$', ln)]
    if bad:
        shown = ', '.join(str(n) for n in bad[:12]) + (' ...' if len(bad) > 12 else '')
        raise SystemExit(
            f"{src_path}: damaged code fence - a line with two backticks where "
            f"a fence needs three, at line(s) {shown}. Restore the third "
            "backtick; nothing was built.")


def collect_headings(lines):
    """(level, title) for every ATX heading in the markdown, skipping fenced
    code blocks so a `#` comment inside a code sample isn't mistaken for one."""
    out = []
    in_code = False
    for ln in lines:
        s = ln.strip()
        if s.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            continue
        if s.startswith('### '):
            out.append((3, _toc_plain(s[4:])))
        elif s.startswith('## '):
            out.append((2, _toc_plain(s[3:])))
        elif s.startswith('# '):
            out.append((1, _toc_plain(s[2:])))
    return out


def add_table_of_contents(doc, headings):
    """A static, pre-populated outline of the document's headings.

    Used when toc_mode == 'static' (the default off Windows). The Windows
    default is add_toc_field(), a real TOC field that Word evaluates.
    Rationale for keeping this path: Word's TOC *field* renders as placeholder
    text until an application re-evaluates it, and on macOS/Linux the
    LibreOffice conversion path can't update it (and LibreOffice's embedded
    Python is launch-constraint-blocked from scripting it). A static outline
    renders identically and correctly in Word, LibreOffice, and any PDF viewer.
    Trade-off: no live page numbers — the section numbers carried in the
    heading text provide the structure instead."""
    p = doc.add_heading("Contents", level=1)
    set_paragraph_shading(p, H1_BG_HEX)

    # Drop a leading level-1 entry (the document title, already on the cover).
    entries = headings[1:] if (headings and headings[0][0] == 1) else headings
    base = min((lv for lv, _ in entries), default=1)
    for level, title in entries:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(18 * (level - base))
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(title)
        run.font.name = FONT_HEAD
        run.font.size = Pt(11)
        if level == base:
            run.font.bold = True

    add_page_break(doc)


def add_banner_paragraph(doc, text):
    """A paragraph with the Heading-1 look (white on the teal banner) but WITHOUT
    the Heading 1 style, so a Word TOC field does not list it. Used for the
    document title in the body and for the "Contents" heading itself."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_HEAD
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = H1_TEXT_RGB
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_paragraph_shading(p, H1_BG_HEX)
    return p


def _ensure_toc_styles(doc):
    """Define Word's built-in 'toc 1..3' paragraph styles in the house look:
    level 1 Myriad bold, lower levels regular, right tab with dot leader at
    the text width so the page numbers line up (DOCUMENT_STANDARDS TOC rule).
    Word applies these when it populates the TOC field. The Hyperlink
    character style is pinned to the house font (the \\h switch makes the
    entries hyperlinks) so no theme font leaks into the PDF."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_TAB_LEADER
    text_width = Inches(8.5 - PAGE_MARGINS['left'] - PAGE_MARGINS['right'])
    for lvl in (1, 2, 3):
        name = f'toc {lvl}'
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        # Word recognises its built-in TOC styles by style ID 'TOC1'..'TOC3'
        # (python-docx derives 'toc1'); with the wrong ID Word treats ours as
        # custom, renames it 'TOC 11', and formats the entries with its own
        # defaults (theme fonts leaked into the PDF, 2026-09-10).
        st.element.set(qn('w:styleId'), f'TOC{lvl}')
        # python-docx marks added styles customStyle=1; a custom style with a
        # built-in name is exactly what makes Word rename it. Drop the flag.
        st.element.attrib.pop(qn('w:customStyle'), None)
        st.base_style = doc.styles['Normal']
        st.font.name = FONT_HEAD
        st.font.size = Pt(10.5 if lvl == 1 else 10)
        st.font.bold = (lvl == 1)
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        pf = st.paragraph_format
        pf.left_indent = Pt(14 * (lvl - 1))
        pf.space_before = Pt(5 if lvl == 1 else 0)
        pf.space_after = Pt(2)
        pf.tab_stops.clear_all()
        pf.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
    try:
        hl = doc.styles['Hyperlink']
    except KeyError:
        hl = doc.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER)
        hl.element.attrib.pop(qn('w:customStyle'), None)
    hl.font.name = FONT_HEAD
    hl.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    hl.font.underline = False


def add_toc_field(doc):
    """Insert a real Word TOC field (levels 1-3, hyperlinked, page numbers).
    It renders as a placeholder line until an application evaluates it:
    the Windows/Word conversion path does that (Fields.Update twice), so the
    PDF carries live page numbers. LibreOffice does NOT evaluate it, which
    is why the static outline (add_table_of_contents) remains the default
    off Windows."""
    add_banner_paragraph(doc, "Contents")
    _ensure_toc_styles(doc)
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z '
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = "Table of contents: open in Word and update fields (F9)."
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)
    add_page_break(doc)


def add_cover_page(doc):
    # Two empty paragraphs at the top push the title down a few cm
    for _ in range(2):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(DOC_TITLE)
    tr.font.name = FONT_HEAD
    tr.font.size = Pt(38)
    tr.font.color.rgb = TITLE_COLOR_RGB
    tr.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run(DOC_SUBTITLE)
    sr.font.name = FONT_HEAD
    sr.font.size = Pt(24)
    sr.font.color.rgb = SUBTITLE_COLOR_RGB
    sr.italic = True

    for _ in range(2):
        doc.add_paragraph()

    # If the project ships a PNG icon, drop it on the cover as a small mark.
    icon_path = Path('icons/dses_workbench.png')
    if icon_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(icon_path), width=Inches(1.5))

    for _ in range(3):
        doc.add_paragraph()

    # Version / author / org / date stacked at the bottom of the cover.
    # DOC_VERSION starting with 'v' renders as "Version X.Y.Z" (manuals);
    # anything else (e.g. "Rev A" for reports) is printed verbatim.
    ver_text = (f"Version {DOC_VERSION.lstrip('v')}"
                if DOC_VERSION.startswith('v') else DOC_VERSION)
    for text in (
        ver_text,
        DOC_AUTHOR,
        DOC_ORG,
        date.today().strftime("%B %Y"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = FONT_HEAD
        r.font.size = Pt(16)
        r.font.color.rgb = SUBTITLE_COLOR_RGB
    add_page_break(doc)


def apply_heading_styles(doc):
    """Tweak the built-in heading styles to match DSES (after the doc is
    created, since python-docx populates styles from the default template)."""
    h1 = doc.styles['Heading 1']
    h1.font.name = FONT_HEAD
    h1.font.color.rgb = H1_TEXT_RGB
    h1.font.size = Pt(11)
    h1.font.bold = True
    # Rick's H1 spacing tweak: 12 pt above, 6 pt below (was 24/0).
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    # Heading-2 stays default (slate); Heading 3 → dark teal
    h2 = doc.styles['Heading 2']
    h2.font.name = FONT_HEAD
    h3 = doc.styles['Heading 3']
    h3.font.name = FONT_HEAD
    h3.font.color.rgb = H3_COLOR_RGB
    h3.font.bold = True

    normal = doc.styles['Normal']
    normal.font.name = FONT_BODY
    normal.font.size = Pt(11)

    # The default template's heading styles name the THEME fonts
    # (asciiTheme="majorHAnsi" = Cambria) and theme attributes take precedence
    # over the explicit face python-docx's font.name sets. Word then copies
    # the heading paragraph-mark formatting into the TOC entries, so Cambria
    # leaked into the PDF's embedded fonts (2026-09-10). Strip the theme
    # attributes and pin every script slot to the house heading face.
    for st in (h1, h2, h3):
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            rf.attrib.pop(qn(f'w:{attr}'), None)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            rf.set(qn(f'w:{attr}'), FONT_HEAD)

    # Also pin the DOCUMENT DEFAULTS (w:docDefaults), not just the Normal
    # style: runless paragraphs (spacers, image anchors, table paragraph
    # marks) fall back to Word's built-in default (Times New Roman / theme
    # Calibri), which then leaks into the PDF's embedded-font list. Same
    # rule as DOCUMENT_STANDARDS.md's generator note, ported to python-docx.
    styles_el = doc.styles.element
    dd = styles_el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults')
        styles_el.insert(0, dd)
    rprd = dd.find(qn('w:rPrDefault'))
    if rprd is None:
        rprd = OxmlElement('w:rPrDefault')
        dd.insert(0, rprd)
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        rprd.append(rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), FONT_BODY)


# ---------------------------------------------------------------------------
# Code blocks + tables (unchanged from previous version)
# ---------------------------------------------------------------------------

def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL_HEX)
    cell.text = ''
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line)
        run.font.name = FONT_MONO
        run.font.size = Pt(9.5)


def add_table(doc, header_row, body_rows, widths_in=None):
    """Markdown table -> Word table. Rows are marked cantSplit (a row never
    breaks across a page) and the header row repeats at the top of every
    page the table spans (Rick, 2026-09-10). `widths_in` (list of inches,
    one per column) fixes the column widths; without it Word autofits."""
    cols = len(header_row)
    # A markdown table whose header cells are all blank ("| | |") is a
    # headerless key/value block (the document-control table): render it
    # without the empty first row (Rick, 2026-09-10).
    headerless = all(not h.strip() for h in header_row)
    if headerless:
        header_row, body_rows = body_rows[0], body_rows[1:]
    table = doc.add_table(rows=1 + len(body_rows), cols=cols)
    table.style = 'Light Grid Accent 1'
    if widths_in:
        table.autofit = False
        for j, w in enumerate(widths_in[:cols]):
            table.columns[j].width = Inches(w)
            for row in table.rows:
                row.cells[j].width = Inches(w)
    for r_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        tr_pr.append(cant)
        if r_idx == 0:
            hdr = OxmlElement('w:tblHeader')
            tr_pr.append(hdr)
    for j, cell_text in enumerate(header_row):
        cell = table.rows[0].cells[j]
        cell.text = ''
        add_runs(cell.paragraphs[0], cell_text)
        if not headerless:
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.name = FONT_HEAD
    for i, row in enumerate(body_rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ''
            cell = table.rows[1 + i].cells[j]
            cell.text = ''
            add_runs(cell.paragraphs[0], cell_text)
    # Word keeps consecutive rows together when every paragraph of a row
    # carries keep-with-next. The header row always keeps with the first body
    # row (no orphaned header at a page foot); a short table keeps whole.
    # Must run AFTER the cell text is written: `cell.text = ...` replaces the
    # cell's paragraphs and would discard the flag (found 2026-09-10).
    n_rows = len(table.rows)
    keep_whole = len(body_rows) <= 10
    for r_idx, row in enumerate(table.rows):
        if r_idx < n_rows - 1 and (r_idx == 0 or keep_whole):
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True


def parse_table_block(lines, start):
    def split_row(s):
        s = s.strip()
        if s.startswith('|'):
            s = s[1:]
        if s.endswith('|'):
            s = s[:-1]
        return [c.strip() for c in s.split('|')]
    i = start
    header = split_row(lines[i]); i += 1
    i += 1  # alignment row
    body = []
    while i < len(lines) and lines[i].strip().startswith('|'):
        body.append(split_row(lines[i]))
        i += 1
    return header, body, i


# ---------------------------------------------------------------------------
# Main parser
# ---------------------------------------------------------------------------

# TOC rendering: 'field' = a real Word TOC field with page numbers (needs the
# Word conversion path to evaluate it); 'static' = pre-populated outline that
# renders anywhere (LibreOffice cannot evaluate the field). Default by platform.
TOC_MODE_DEFAULT = 'field' if sys.platform == 'win32' else 'static'


def md_to_docx(src_path: Path, dst_path: Path, cover: bool = True,
               toc: bool = True, toc_mode: str = None):
    toc_mode = toc_mode or TOC_MODE_DEFAULT
    text = src_path.read_text(encoding='utf-8')
    check_fences(text, src_path)
    text = smartify_quotes(text)
    lines = text.split('\n')

    doc = Document()
    configure_section(doc.sections[0])
    apply_heading_styles(doc)
    if cover:
        add_cover_page(doc)
    if toc and toc_mode == 'field':
        add_toc_field(doc)
    elif toc:
        add_table_of_contents(doc, collect_headings(lines))

    # Style a Heading-1 paragraph so it gets the teal banner. We do this by
    # post-processing each Heading-1 paragraph after add_heading() rather
    # than wiring it into the global style (which would also paint the TOC,
    # the cover page, etc., with shading).
    # The document's first level-1 heading is its title (already on the
    # cover): render it as a banner paragraph WITHOUT the Heading 1 style so
    # a TOC field does not list it (the static outline drops it likewise).
    seen_h1 = {'n': 0}

    def add_h1_banner(text):
        seen_h1['n'] += 1
        if cover and seen_h1['n'] == 1:
            add_banner_paragraph(doc, text)
            return
        p = doc.add_heading(text, level=1)
        set_paragraph_shading(p, H1_BG_HEX)

    i = 0
    prev_was_table = False          # for post-table paragraph spacing
    pending_widths = None           # from a '<!-- widths: ... -->' comment
    open_num = None                 # numbered list a later item may continue
    in_list = False                 # previous block was a list item or its continuation
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        # A body paragraph immediately under a table gets a little air above
        # it (Rick's house tweak). Latch the "previous block was a table"
        # state here and clear it each iteration; only the table branch re-sets
        # it, so a heading/image between table and paragraph correctly resets.
        after_table = prev_was_table
        prev_was_table = False

        if stripped.startswith('```'):
            # A fence indented under a list item: drop that indent from the
            # code lines, keeping any deeper, relative indentation.
            pad = len(line) - len(line.lstrip(' '))
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(_dedent(lines[i], pad))
                i += 1
            i += 1
            add_code_block(doc, code_lines)
            continue

        # Single-line HTML comments are authoring notes, except the table
        # width hint '<!-- widths: 1.2,3.0,2.3 -->' (inches per column) that
        # applies to the next table.
        m_c = re.match(r'^<!--\s*(.*?)\s*-->$', stripped)
        if m_c:
            m_w = re.match(r'widths:\s*([\d.,\s]+)$', m_c.group(1))
            if m_w:
                pending_widths = [float(x) for x in m_w.group(1).split(',')
                                  if x.strip()]
            i += 1; continue

        if stripped.startswith('|') and i + 1 < len(lines) \
                and re.match(r'^\s*\|[\s\-:|]+\|\s*$', lines[i + 1]):
            header, body, i = parse_table_block(lines, i)
            add_table(doc, header, body, widths_in=pending_widths)
            pending_widths = None
            prev_was_table = True
            continue

        # Headings print as plain text: their inline markers (code, bold,
        # links) are dropped, as the table of contents already did. A heading
        # also closes any numbered list.
        if stripped.startswith('### '):
            doc.add_heading(_toc_plain(stripped[4:]), level=3)
            open_num, in_list = None, False
            i += 1; continue
        if stripped.startswith('## '):
            doc.add_heading(_toc_plain(stripped[3:]), level=2)
            open_num, in_list = None, False
            i += 1; continue
        if stripped.startswith('# '):
            add_h1_banner(_toc_plain(stripped[2:]))
            open_num, in_list = None, False
            i += 1; continue

        # Image: a line of the form ![caption](path). The path is resolved
        # relative to the source .md; the alt text becomes an italic caption.
        m_img = re.match(r'^!\[(.*?)\]\((.+?)\)\s*$', stripped)
        if m_img:
            caption, img_ref = m_img.group(1), m_img.group(2)
            img_path = Path(img_ref)
            if not img_path.is_absolute():
                img_path = src_path.parent / img_path
            if img_path.is_file():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img_path), width=Inches(6.2))
                if caption:
                    cp = doc.add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr = cp.add_run(caption)
                    cr.font.size = Pt(8.5)
                    cr.font.italic = True
                    cr.font.color.rgb = SUBTITLE_COLOR_RGB
            else:
                p = doc.add_paragraph()
                add_runs(p, f"*[missing image: {img_ref}]*")
            i += 1; continue

        if re.match(r'^\s*[-*]\s+', line):
            while i < len(lines) and re.match(r'^\s*[-*]\s+', lines[i]):
                content = re.sub(r'^\s*[-*]\s+', '', lines[i])
                i += 1
                # Coalesce hard-wrapped continuation lines (indented, not a
                # new bullet / number / blank) into the SAME bullet — else
                # they render as detached body paragraphs after the item.
                while (i < len(lines) and lines[i].startswith('  ')
                       and lines[i].strip() != ''
                       and not re.match(r'^\s*[-*]\s+', lines[i])
                       and not re.match(r'^\s*\d+\.\s+', lines[i])
                       and not lines[i].strip().startswith('```')):
                    content += ' ' + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style='List Bullet')
                add_runs(p, content)
            in_list = True
            continue

        if re.match(r'^\s*\d+\.\s+', line):
            # "1." starts a new list and the numbering restarts. Any other
            # number continues the list that a code block, a nested bullet
            # list or an indented paragraph interrupted; without this, step 3
            # of a procedure printed as "1." (27-Sep to 4-Oct-2026).
            first_no = int(re.match(r'^\s*(\d+)\.', line).group(1))
            if first_no == 1 or open_num is None:
                open_num = new_numbered_list(doc)
            list_num = open_num
            in_list = True
            while i < len(lines) and re.match(r'^\s*\d+\.\s+', lines[i]):
                content = re.sub(r'^\s*\d+\.\s+', '', lines[i])
                p = doc.add_paragraph(style='List Number')
                set_list_num(p, list_num)
                add_runs(p, content)
                while (i + 1 < len(lines) and lines[i + 1].startswith('   ')
                       and not re.match(r'^\s*\d+\.\s+', lines[i + 1])
                       and not re.match(r'^\s*[-*]\s+', lines[i + 1])
                       and lines[i + 1].strip() != ''):
                    i += 1
                    cont = lines[i].strip()
                    if cont.startswith('```'):
                        pad = len(lines[i]) - len(lines[i].lstrip(' '))
                        code_lines = []
                        i += 1
                        while i < len(lines) and not lines[i].strip().startswith('```'):
                            code_lines.append(_dedent(lines[i], pad))
                            i += 1
                        add_code_block(doc, code_lines)
                    elif cont:
                        cp = doc.add_paragraph()
                        cp.paragraph_format.left_indent = Inches(0.5)
                        add_runs(cp, cont)
                i += 1
            continue

        if stripped == '':
            prev_was_table = after_table   # carry table-state across blank lines
            i += 1; continue

        # Paragraph: coalesce wrapped lines
        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if ns == '' or ns.startswith(('#', '-', '*', '|', '```')) \
                    or re.match(r'^\s*\d+\.\s+', nxt):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        if after_table:
            p.paragraph_format.space_before = Pt(6)
        # A paragraph indented under a list item belongs to that item: keep
        # it aligned with the item's text. An unindented paragraph ends the
        # list context (a numbered list may still be continued by number).
        if in_list and line.startswith('  '):
            p.paragraph_format.left_indent = Inches(0.5)
        else:
            in_list = False
        para_text = ' '.join(s.strip() for s in para_lines)
        add_runs(p, para_text)
        # A lead-in paragraph ("The decisions are:") stays with the list,
        # table, or code block it introduces instead of hanging at the foot
        # of a page (Rick, 2026-09-10).
        if para_text.rstrip().endswith(':'):
            p.paragraph_format.keep_with_next = True

    doc.save(dst_path)
    print(f"Wrote {dst_path} ({dst_path.stat().st_size // 1024} KB)")


def _kill_stale_invisible_word():
    """Kill orphaned background Word processes left over from a previous
    run that died between Documents.Open and word.Quit. A Word that owns a
    VISIBLE window is someone's open document and is left alone (added
    27-Sep-2026; before that every winword.exe was killed). Skipped
    silently if psutil isn't available."""
    try:
        import psutil
    except ImportError:
        return
    visible = set()
    try:
        import win32gui
        import win32process

        def _cb(hwnd, acc):
            if win32gui.IsWindowVisible(hwnd):
                acc.add(win32process.GetWindowThreadProcessId(hwnd)[1])
        win32gui.EnumWindows(_cb, visible)
    except Exception:
        pass
    for p in psutil.process_iter(['name', 'pid']):
        if (p.info.get('name') or '').lower() == 'winword.exe':
            if p.info.get('pid') in visible:
                continue
            try:
                p.kill()
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass


def _find_soffice():
    """Locate a LibreOffice/OpenOffice headless binary, or None."""
    import shutil
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    # Common macOS install location (not on PATH by default).
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.is_file():
        return str(mac)
    return None


def _convert_docx_to_pdf_soffice(soffice: str, docx_path: Path, pdf_path: Path):
    """macOS/Linux PDF path: convert via LibreOffice headless.

    Note: LibreOffice does not re-evaluate Word's TOC *field* the way Word
    does, so the table of contents may render with its placeholder text. The
    body, headings, page numbers and styling come through fine. For a fully
    populated TOC, build on Windows with Word (the path below) or open the
    .docx in Word once and Save As PDF."""
    import subprocess
    outdir = pdf_path.resolve().parent
    subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(docx_path.resolve())],
                   check=True, capture_output=True, text=True)
    produced = outdir / (docx_path.stem + ".pdf")
    if produced != pdf_path.resolve():
        produced.replace(pdf_path)
    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size // 1024} KB) via LibreOffice")


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path):
    """Convert the DOCX to PDF using the best available engine:

      * Windows + Microsoft Word -> Word via pywin32 (updates the TOC/PAGE
        fields properly), then a PDF printer driver chosen by the
        DSES_PDF_ENGINE environment variable (see _pdf_engine);
      * otherwise (macOS/Linux)  -> LibreOffice headless, if installed.

    Raises if neither is available so the caller can fall back to shipping
    the .docx and asking for a manual Save-As-PDF."""
    if sys.platform == "win32":
        return _convert_docx_to_pdf_word(docx_path, pdf_path)

    soffice = _find_soffice()
    if soffice:
        return _convert_docx_to_pdf_soffice(soffice, docx_path, pdf_path)
    raise RuntimeError(
        "No PDF engine found. On Windows install Microsoft Word; on "
        "macOS/Linux install LibreOffice (macOS: `brew install --cask "
        "libreoffice`) and re-run, or open the .docx and Save As PDF.")


def _printer_available(name):
    """True if a printer of that name is installed for this user."""
    try:
        import win32print
        names = {p[2] for p in win32print.EnumPrinters(
            win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS)}
        return name in names
    except Exception:
        return False


def _adobe_pdf_printer_available():
    """True if the 'Adobe PDF' printer (Acrobat Distiller) is installed.

    Note this is independent of Acrobat.exe itself: on the Windows dev box
    Acrobat has been crashing since its 2026-08-02 update, but the printer
    driver and Distiller are separate binaries (verified 19-Aug-2026) —
    although the driver DOES launch Acrobat.exe for font capture, see
    _pdf_engine."""
    return _printer_available("Adobe PDF")


def _find_acrodist():
    """Full path of Acrobat Distiller's acrodist.exe, or None."""
    for c in (Path(r"C:/Program Files (x86)/Adobe/Acrobat DC/Acrobat/acrodist.exe"),
              Path(r"C:/Program Files/Adobe/Acrobat DC/Acrobat/acrodist.exe")):
        if c.is_file():
            return c
    return None


def _word_printing(word):
    """True while Word still has a background print job in flight."""
    try:
        return int(word.BackgroundPrintingStatus) > 0
    except Exception:
        return False


def _newest_winword_pid(since):
    """PID of the Word instance started after `since` (ours), or None."""
    try:
        import psutil
    except ImportError:
        return None
    best = None
    for p in psutil.process_iter(['name', 'pid', 'create_time']):
        if (p.info.get('name') or '').lower() != 'winword.exe':
            continue
        if (p.info.get('create_time') or 0) >= since - 5:
            if best is None or p.info['create_time'] > best[0]:
                best = (p.info['create_time'], p.info['pid'])
    return best[1] if best else None


def _acrobat_helpers_since(t0):
    """Acrobat.exe processes launched after t0 — the Adobe PDF driver starts
    one for font capture — as (age_seconds, psutil.Process) pairs."""
    try:
        import psutil
    except ImportError:
        return []
    out = []
    now = __import__("time").time()
    for p in psutil.process_iter(['name', 'create_time']):
        if (p.info.get('name') or '').lower() == 'acrobat.exe':
            ct = p.info.get('create_time') or 0
            if ct >= t0 - 1:
                out.append((now - ct, p))
    return out


def _kill_acrobat_helpers(t0):
    """Kill the font-capture Acrobat.exe processes started after t0 (a
    crashed one sits in its 'Font Capture' error box until killed). Returns
    the pids killed."""
    killed = []
    for _, p in _acrobat_helpers_since(t0):
        try:
            pid = p.pid
            p.kill()
            killed.append(str(pid))
        except Exception:
            pass
    return killed


def _cancel_print_jobs(printer_name):
    """Delete every job in the named printer's queue (best effort)."""
    try:
        import win32print
        h = win32print.OpenPrinter(printer_name)
        try:
            for job in win32print.EnumJobs(h, 0, 999, 1):
                try:
                    win32print.SetJob(h, job['JobId'], 0, None,
                                      win32print.JOB_CONTROL_DELETE)
                except Exception:
                    pass
        finally:
            win32print.ClosePrinter(h)
    except Exception:
        pass


def _print_to_adobe_pdf(word, doc, docx_path: Path, pdf_path: Path,
                        ps_timeout=120.0, helper_grace=15.0):
    """Produce the PDF via Word -> PostScript file -> Distiller, directly.

    Word's own exporter (SaveAs FileFormat=17) cannot embed OpenType-PS (CFF)
    faces and silently rasterizes every run that uses one — with the house
    fonts that means the whole document loses its text layer, so the PDF must
    come from Adobe's pipeline instead (DOCUMENT_STANDARDS.md section 8).

    We deliberately do NOT print through the "Adobe PDF" printer *port*: its
    port monitor drops the PDF wherever its "last used folder" points, keeps
    the file locked long afterwards, and when a print collides with such a
    leftover it jams the queue with stacked modal error dialogs (19-Aug-2026
    incident). Instead, Word prints PostScript to a scratch FILE we name
    (PrintToFile), and acrodist.exe distills that file synchronously — fully
    deterministic paths, no spooler, no dialogs.

    BOUNDED since 27-Sep-2026: the print runs in Word's background so this
    process keeps control. The Adobe driver launches Acrobat.exe for "font
    capture" when a document uses a glyph outside its fonts (the install
    guide's two symbol characters); a broken Acrobat.exe then dies into a
    modal error box and the job never produces PostScript. If such a helper
    has been alive for `helper_grace` seconds with no PostScript yet, or
    nothing has arrived within `ps_timeout`, the helper is killed, the job
    cancelled, and a RuntimeError sends the caller to the msprint fallback.
    A synchronous PrintOut sat on that box until a human clicked OK."""
    import subprocess
    import shutil
    import time

    acrodist = _find_acrodist()
    if acrodist is None:
        raise RuntimeError("acrodist.exe not found under Program Files — "
                           "is Acrobat/Distiller installed?")

    scratch_dir = Path(tempfile.mkdtemp(prefix="dses_distill_"))
    ps_file = scratch_dir / (pdf_path.stem + ".ps")
    out_pdf = scratch_dir / (pdf_path.stem + ".pdf")

    previous_printer = word.ActivePrinter
    t0 = time.time()
    killed = []
    aborted = None
    completed = False
    # A hung Distiller or a stuck AcroTray blocks the Adobe PDF driver, and
    # then PrintOut never returns (8-Oct-2026): clear both before printing.
    kill_stale_distillers("before printing")
    kill_acrotray("before printing")

    # PrintOut(Background=True) is supposed to return at once, but when the
    # driver is blocked it does not return at all, so the watchdog below is
    # a thread: if no PostScript has appeared `stall_grace` seconds after
    # the call it kills AcroTray (the proven remedy) and any hung Distiller,
    # and says so; the main loop's own clock starts only when PrintOut
    # returns, so a remedied stall is not then mistaken for a timeout.
    import threading
    stall_grace = 30.0
    returned = threading.Event()
    stop_watch = threading.Event()
    watch_notes = []

    def _watchdog():
        t_call = time.time()
        remedied = False
        while not stop_watch.wait(1.0):
            if ps_file.is_file() and ps_file.stat().st_size > 0:
                return
            if not remedied and time.time() - t_call > stall_grace:
                remedied = True
                k = kill_acrotray("after %.0f s with no PostScript (driver stalled)"
                                  % (time.time() - t_call))
                k += kill_stale_distillers("after the stall")
                watch_notes.append("stall remedy at %.0f s: killed pids %s"
                                   % (time.time() - t_call, k or "none"))
    watch = threading.Thread(target=_watchdog, daemon=True,
                             name="adobe-pdf-stall-watchdog")
    try:
        word.ActivePrinter = "Adobe PDF"
        watch.start()
        doc.PrintOut(Background=True, PrintToFile=True,
                     OutputFileName=str(ps_file))
        returned.set()
        t_ret = time.time()
        if t_ret - t0 > 5:
            print("build_doc: PrintOut blocked for %.0f s before returning%s"
                  % (t_ret - t0, ("; " + "; ".join(watch_notes)) if watch_notes else ""),
                  file=sys.stderr)
        last, stable = -1, 0
        deadline = t_ret + ps_timeout
        last_growth = t_ret
        abort_by = None
        while True:
            size = ps_file.stat().st_size if ps_file.is_file() else -1
            busy = _word_printing(word)
            if size > 0 and size == last:
                stable += 1
            else:
                stable = 0
                if size > 0:
                    last_growth = time.time()
            if size > 0 and stable >= 2 and not busy and abort_by is None:
                completed = True
                break
            now = time.time()
            helper_age = max([a for a, _ in _acrobat_helpers_since(t0)] or [0.0])
            # Abort only when nothing is being written: no PostScript past
            # the deadline, a growing file that stopped growing for 60 s,
            # or a crashed font-capture helper with no PostScript.
            stalled = (size <= 0 and now > deadline) or \
                      (size > 0 and now - last_growth > 60.0 and busy)
            if abort_by is None and (stalled or
                                     (size <= 0 and helper_age > helper_grace)):
                aborted = ("Acrobat.exe font-capture helper alive %.0f s with "
                           "no PostScript" % helper_age
                           if (size <= 0 and helper_age > helper_grace) else
                           "no PostScript within %.0f s of PrintOut returning"
                           % ps_timeout if size <= 0 else
                           "PostScript stopped growing at %d bytes" % size)
                killed = _kill_acrobat_helpers(t0)
                _cancel_print_jobs("Adobe PDF")
                abort_by = now + 20.0
            if abort_by is not None and (now > abort_by or not busy):
                break
            last = size
            time.sleep(1.0)
    finally:
        stop_watch.set()
        try:
            word.ActivePrinter = previous_printer
        except Exception:
            pass
    if not completed:
        detail = aborted or "print did not finish"
        if killed:
            detail += " — killed Acrobat.exe pid " + ", ".join(killed)
        shutil.rmtree(scratch_dir, ignore_errors=True)
        raise RuntimeError(f"Word wrote no usable PostScript ({detail})")

    run_distiller(acrodist, ps_file, out_pdf)

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    shutil.copy2(str(out_pdf), str(pdf_path))
    shutil.rmtree(scratch_dir, ignore_errors=True)


# ---------------------------------------------------------------------------
# Distiller runner and cross-session serialization (8-Oct-2026)
#
# Two things went wrong on 7 and 8 October 2026 and both are fixed here.
#
# 1. acrodist.exe /N /Q writes the complete PDF in about a second and then,
#    some of the time, hangs on exit ("Acrobat Distiller (Not Responding)",
#    End of Job in its log pane). The old subprocess.run(..., timeout=300)
#    waited on the process, not the PDF: five minutes of nothing, then a
#    TimeoutExpired that sent a perfectly good build to the msprint fallback
#    (fonts embedded whole, bigger files) — which Rick does not want. And a
#    Distiller left hanging blocks the Adobe PDF driver for every later
#    print on the machine, so Word's PrintOut then produces no PostScript at
#    all (the 7-Oct stalls: ~20 empty dses_distill_* folders in %TEMP%).
#    run_distiller() therefore watches the OUTPUT (the PDF exists, has
#    stopped growing, ends in %%EOF and parses), and once it is complete it
#    kills a Distiller that has not quit by itself. Hung or leftover
#    Distillers are killed before every print.
#
# 2. Two build_doc.py builds in two Claude sessions at the same time killed
#    each other's invisible Word ("The RPC server is unavailable"), because
#    _kill_stale_invisible_word() cannot tell an orphan from another
#    session's live build. build_lock is a machine-wide named mutex held
#    for the whole Word + print step: a second build WAITS (saying so on
#    stderr) instead of colliding; a build that dies releases it (the OS
#    marks the mutex abandoned and the next waiter gets it). Under the lock,
#    any invisible Word or Distiller that is still alive IS an orphan, so
#    the kills become safe. pptx_to_pdf.py (the deck route) uses both.
# ---------------------------------------------------------------------------

BUILD_MUTEX_NAME = r"Global\DSES_build_doc_word_pdf"


class build_lock:
    """Context manager: one Word + PDF-printer build at a time on this
    machine, across processes and Claude sessions (a named Windows mutex;
    a no-op elsewhere). Waits up to `wait_s` for another build to finish,
    reporting the wait on stderr; a crashed holder releases it (abandoned
    mutex = acquired)."""

    def __init__(self, what="build", wait_s=1800.0):
        self.what = what
        self.wait_s = wait_s
        self._h = None
        self.held = False

    def __enter__(self):
        if sys.platform != "win32":
            return self
        import time
        import win32event
        self._h = win32event.CreateMutex(None, False, BUILD_MUTEX_NAME)
        t0 = time.time()
        said = False
        while True:
            rc = win32event.WaitForSingleObject(self._h, 5000)
            if rc in (win32event.WAIT_OBJECT_0, win32event.WAIT_ABANDONED):
                break
            if rc != win32event.WAIT_TIMEOUT:
                raise RuntimeError(f"build lock wait failed (code {rc})")
            if not said:
                print(f"build_doc: another Word/PDF build holds the machine "
                      f"lock; waiting for it before {self.what} "
                      f"(up to {self.wait_s / 60:.0f} min)", file=sys.stderr)
                said = True
            if time.time() - t0 > self.wait_s:
                raise RuntimeError("gave up waiting for the other Word/PDF "
                                   f"build after {self.wait_s / 60:.0f} min")
        if said:
            print(f"build_doc: lock acquired after {time.time() - t0:.0f} s",
                  file=sys.stderr)
        self.held = True
        return self

    def __exit__(self, *exc):
        if self._h is not None:
            import win32event
            import win32api
            if self.held:
                try:
                    win32event.ReleaseMutex(self._h)
                except Exception:
                    pass
            try:
                win32api.CloseHandle(self._h)
            except Exception:
                pass
            self._h = None
            self.held = False
        return False


def _hung_windows_of(pid):
    """True if the process owns a top-level window Windows reports as not
    responding (what the title bar shows as '(Not Responding)')."""
    try:
        import win32gui
        import win32process
    except ImportError:
        return False
    hung = []

    def _cb(hwnd, acc):
        try:
            if win32process.GetWindowThreadProcessId(hwnd)[1] == pid \
                    and win32gui.IsHungAppWindow(hwnd):
                acc.append(hwnd)
        except Exception:
            pass
    try:
        win32gui.EnumWindows(_cb, hung)
    except Exception:
        pass
    return bool(hung)


def kill_stale_distillers(when="before printing"):
    """Kill acrodist.exe processes that are ours (command line names a
    dses_distill_ scratch file) or that are not responding, and report each
    on stderr. A Distiller Rick opened himself and that is answering is left
    alone. Returns the pids killed."""
    try:
        import psutil
    except ImportError:
        return []
    killed = []
    for p in psutil.process_iter(['name', 'pid', 'cmdline', 'create_time']):
        if (p.info.get('name') or '').lower() != 'acrodist.exe':
            continue
        cmd = " ".join(p.info.get('cmdline') or [])
        ours = "dses_distill_" in cmd
        hung = _hung_windows_of(p.info['pid'])
        if not (ours or hung):
            continue
        try:
            p.kill()
            p.wait(timeout=10)
            killed.append(p.info['pid'])
            print(f"build_doc: killed a leftover Distiller (pid {p.info['pid']}, "
                  f"{'ours' if ours else 'not responding'}) {when}",
                  file=sys.stderr)
        except Exception:
            pass
    return killed


def kill_acrotray(when="before printing"):
    """Kill Acrobat's tray helper, AcroTray.exe. The Adobe PDF printer
    driver hands every job through it, and one whose Distiller child was
    killed (the hang-on-exit case) sits there answering but never passing
    the next job on: Word's PrintOut then blocks with no PostScript for as
    long as you care to wait (7-Oct and 8-Oct-2026; proven on 8-Oct by
    killing it mid-stall — PostScript appeared five seconds later). The
    driver relaunches a fresh one on the next print, so killing it costs
    nothing. Returns the pids killed."""
    try:
        import psutil
    except ImportError:
        return []
    killed = []
    for p in psutil.process_iter(['name', 'pid']):
        if (p.info.get('name') or '').lower() != 'acrotray.exe':
            continue
        try:
            p.kill()
            p.wait(timeout=10)
            killed.append(p.info['pid'])
            print(f"build_doc: killed AcroTray.exe (pid {p.info['pid']}) {when}",
                  file=sys.stderr)
        except Exception:
            pass
    return killed


def _pdf_complete(path: Path) -> bool:
    """True once the file ends in a %%EOF trailer and parses with at least
    one page (pypdf when available; the trailer check alone otherwise)."""
    try:
        with open(path, "rb") as fh:
            fh.seek(max(0, path.stat().st_size - 1024))
            tail = fh.read()
    except OSError:
        return False
    if b"%%EOF" not in tail:
        return False
    try:
        from pypdf import PdfReader
    except ImportError:
        return True
    try:
        return len(PdfReader(str(path)).pages) >= 1
    except Exception:
        return False


def run_distiller(acrodist, ps_file: Path, out_pdf: Path,
                  timeout=300.0, exit_grace=5.0):
    """Distill `ps_file` to `out_pdf` (Distiller writes it beside the .ps)
    and return when the PDF is complete, whether or not acrodist.exe has
    quit. Completion = the PDF exists, has stopped growing, ends in %%EOF
    and parses. A Distiller that is still alive `exit_grace` seconds after
    that is killed (the 8-Oct-2026 hang-on-exit); one that writes no
    complete PDF within `timeout` is killed and a RuntimeError raised with
    the tail of its log. Never waits on process exit, never captures its
    pipes (a hung GUI process would hold them open)."""
    import subprocess
    import time

    kill_stale_distillers("before distilling")
    log = ps_file.with_suffix(".log")
    for stale in (out_pdf, log):
        try:
            stale.unlink()
        except OSError:
            pass
    p = subprocess.Popen([str(acrodist), "/N", "/Q", str(ps_file)],
                         stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                         stderr=subprocess.DEVNULL)
    t0 = time.time()
    last, stable, complete = -1, 0, False
    while time.time() - t0 < timeout:
        rc = p.poll()
        size = out_pdf.stat().st_size if out_pdf.is_file() else -1
        stable = stable + 1 if (size > 0 and size == last) else 0
        last = size
        if size > 0 and (stable >= 2 or rc is not None) and _pdf_complete(out_pdf):
            complete = True
            break
        if rc is not None and size <= 0:
            break                       # exited without writing anything
        time.sleep(0.5)
    if p.poll() is None:
        if complete:
            try:
                p.wait(timeout=exit_grace)
            except subprocess.TimeoutExpired:
                pass
        if p.poll() is None:
            p.kill()
            try:
                p.wait(timeout=10)
            except subprocess.TimeoutExpired:
                pass
            print("build_doc: Distiller (pid %d) %s; killed it" % (
                p.pid,
                "hung on exit after writing the PDF in %.0f s" % (time.time() - t0)
                if complete else
                "wrote no complete PDF within %.0f s" % timeout), file=sys.stderr)
    if not complete:
        detail = log.read_text(errors="replace")[-2000:] if log.is_file() else "(no log)"
        raise RuntimeError(f"Distiller produced no PDF from {ps_file}: {detail}")
    return out_pdf


def _cleanup_scratch(max_age_s=86400.0):
    """Remove dses_distill_* scratch folders left in %TEMP% by earlier runs:
    empty ones, and any older than a day."""
    import shutil
    import time
    now = time.time()
    try:
        for d in Path(tempfile.gettempdir()).glob("dses_distill_*"):
            try:
                if not d.is_dir():
                    continue
                if not any(d.iterdir()) or now - d.stat().st_mtime > max_age_s:
                    shutil.rmtree(d, ignore_errors=True)
            except OSError:
                pass
    except OSError:
        pass


def _remove_stale_owner_file(docx_path: Path):
    """Delete Word's '~$name.docx' owner file if one is left beside the
    document and no Word is running (a build killed mid-print leaves it,
    and Word then opens the document read-only without saying so)."""
    owner = docx_path.with_name("~$" + docx_path.name[2:])
    if not owner.exists():
        return
    try:
        import psutil
        if any((p.info.get('name') or '').lower() == 'winword.exe'
               for p in psutil.process_iter(['name'])):
            return
        owner.unlink()
        print(f"build_doc: removed stale Word owner file {owner.name}",
              file=sys.stderr)
    except Exception:
        pass


def _print_to_ms_pdf(word, doc, pdf_path: Path):
    """Fallback PDF route: print through 'Microsoft Print to PDF' straight
    to a file. Verified 2026-09-18 on the Board priorities document: the
    OpenType-PS house fonts come through embedded (as CID TrueType), text is
    selectable, 39 pp in 0.8 MB. Not as clean as Distiller (fonts are
    renamed CIDFont+Fn) but far better than Word's exporter, which
    rasterizes them. Restores the previous active printer afterwards."""
    import time
    if not _printer_available("Microsoft Print to PDF"):
        raise RuntimeError("the 'Microsoft Print to PDF' printer is not "
                           "installed (Windows optional feature)")
    out = str(pdf_path.resolve())
    try:
        Path(out).unlink()
    except OSError:
        pass
    prev = word.ActivePrinter
    try:
        word.ActivePrinter = "Microsoft Print to PDF"
        doc.PrintOut(Background=False, PrintToFile=True, OutputFileName=out)
        # The spooler can finish a moment after PrintOut returns.
        for _ in range(60):
            if Path(out).exists() and Path(out).stat().st_size > 0:
                break
            time.sleep(0.5)
        else:
            raise RuntimeError("Microsoft Print to PDF produced no file")
    finally:
        try:
            word.ActivePrinter = prev
        except Exception:
            pass


def _dismiss_adobe_pdf_dialog():
    """Close the modal 'Adobe PDF' error box the printer driver leaves on
    screen when it refuses to make PostScript (the 'rely on system fonts'
    trap). Best effort; silent if there is none or win32gui is missing."""
    try:
        import win32gui, win32con
    except ImportError:
        return
    def _cb(hwnd, found):
        if win32gui.IsWindowVisible(hwnd) and win32gui.GetWindowText(hwnd) == "Adobe PDF":
            found.append(hwnd)
    found = []
    try:
        win32gui.EnumWindows(_cb, found)
        for hwnd in found:
            win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
    except Exception:
        pass


PDF_ENGINES = ("auto", "msprint", "distiller", "pdfmaker")


def _pdf_engine():
    """Which Windows PDF engine to use: the DSES_PDF_ENGINE environment
    variable, one of distiller | msprint | pdfmaker | auto (default).

    auto = distiller when the 'Adobe PDF' printer and acrodist.exe are
    installed, else msprint. The Distiller attempt is BOUNDED and falls back
    to msprint by itself (see _print_to_adobe_pdf), so the daily output —
    named MinionPro/MyriadPro subsets — is produced whenever Acrobat
    cooperates, and the build never waits on a human when it does not.

    History (27-Sep-2026): the Adobe PDF driver and the PDFMaker add-in
    both launch Acrobat.exe for "font capture" when a document uses a glyph
    outside its fonts (the install guide's two symbol characters); that
    day Acrobat.exe died into a modal "Font Capture: Windows - Application
    Error" box (0xc06d007e) that only a human click dismissed — Distiller
    sat on it, PDFMaker hung Word on it — and 1.6.0 was cut with msprint
    as the default. Rick: "Distiller first with the fallback"; the bounded
    attempt replaced that the same evening."""
    name = (os.environ.get("DSES_PDF_ENGINE") or "auto").strip().lower()
    if name not in PDF_ENGINES:
        raise RuntimeError(f"DSES_PDF_ENGINE={name!r}: expected one of "
                           f"{', '.join(PDF_ENGINES)}")
    if name == "auto":
        if _adobe_pdf_printer_available() and _find_acrodist() is not None:
            return "distiller"
        return "msprint"
    return name


def _word_app():
    """An EARLY-BOUND Word.Application (pywin32 makepy wrapper), or raise.

    Early binding is not optional: with a plain dynamic Dispatch the NAMED
    arguments of PrintOut silently misbind, PrintToFile/OutputFileName are
    lost, and Word puts up 'Save Print Output As' for someone to cancel
    (27-Sep-2026; the 19-Aug-2026 variant wrote nothing at all). The usual
    reason EnsureDispatch fails is a stale makepy cache after an Office
    update (AttributeError ... CLSIDToClassMap / CLSIDToPackageMap): that
    cache is purged and the binding retried once. Still failing -> raise;
    this function never degrades to Dispatch."""
    import shutil
    import win32com
    from win32com.client import gencache
    last = None
    for attempt in (1, 2):
        try:
            return gencache.EnsureDispatch('Word.Application')
        except AttributeError as exc:
            last = exc
            if attempt == 2 or "CLSIDTo" not in str(exc):
                break
            cache = Path(win32com.__gen_path__)
            print(f"build_doc: stale pywin32 makepy cache ({exc}); purging "
                  f"{cache} and retrying", file=sys.stderr)
            shutil.rmtree(cache, ignore_errors=True)
            for name in [k for k in sys.modules if k.startswith("win32com.gen_py")]:
                del sys.modules[name]
            try:
                gencache.GetGeneratePath()
                gencache.Rebuild(verbose=0)
            except Exception:
                pass
    raise RuntimeError(
        f"Word could not be early-bound ({last}). Delete the pywin32 cache "
        "folder %LOCALAPPDATA%/Temp/gen_py and rerun; a dynamic Dispatch is "
        "refused because its PrintOut would stop on a print dialog.")


class _DialogSentinel:
    """Background thread that keeps a scripted print from parking on a
    modal box nobody is there to click: print-to-file prompts get WM_CLOSE
    (= Cancel, so the build fails instead of waiting), any dialog owned by
    OUR invisible Word instance (the 'Adobe PDF' printer-setup box, a
    'Microsoft Word' alert) is closed the same way, and the Acrobat 'Font
    Capture' crash box gets its OK button when it enumerates (it did not
    always on 27-Sep-2026 — _print_to_adobe_pdf's helper kill covers that).
    Everything it touches is reported on stderr."""
    CANCEL = {"Save Print Output As", "Save PDF File As", "Adobe PDF"}

    def __init__(self, word_pid=None):
        import threading
        self.word_pid = word_pid
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._run, daemon=True,
                                        name="pdf-dialog-sentinel")
        self.handled = []

    def __enter__(self):
        self._thread.start()
        return self

    def __exit__(self, *exc):
        self._stop.set()
        self._thread.join(timeout=3)
        return False

    def _run(self):
        try:
            import win32con
            import win32gui
            import win32process
        except ImportError:
            return

        def scan(hwnd, acc):
            if not win32gui.IsWindowVisible(hwnd):
                return
            title = win32gui.GetWindowText(hwnd)
            if title in self.CANCEL or title.startswith("Font Capture"):
                acc.append((hwnd, title))
            elif self.word_pid and win32gui.GetClassName(hwnd) == "#32770":
                try:
                    if win32process.GetWindowThreadProcessId(hwnd)[1] == self.word_pid:
                        acc.append((hwnd, title or "(untitled Word dialog)"))
                except Exception:
                    pass

        def press_ok(hwnd):
            buttons = []

            def child(h, acc):
                if (win32gui.GetClassName(h) == "Button"
                        and win32gui.GetWindowText(h).replace("&", "") == "OK"):
                    acc.append(h)
            win32gui.EnumChildWindows(hwnd, child, buttons)
            for b in buttons:
                win32gui.SendMessage(b, win32con.BM_CLICK, 0, 0)
            return bool(buttons)

        # A dialog owned by our Word that is NOT a known print prompt (for
        # example the Adobe driver's 'Create Adobe PDF' box seen 8-Oct-2026
        # during a stalled print) is given `grace` seconds on screen before
        # it is closed, in case it is a progress window of a job that is
        # still working; the known prompts and the Font Capture crash box
        # are closed at once.
        grace = 90.0
        first_seen = {}
        while not self._stop.is_set():
            found = []
            try:
                win32gui.EnumWindows(scan, found)
                now = __import__("time").time()
                keep = []
                for hwnd, title in found:
                    known = title in self.CANCEL or title.startswith("Font Capture")
                    if not known:
                        first_seen.setdefault(hwnd, now)
                        if now - first_seen[hwnd] < grace:
                            if first_seen[hwnd] == now:
                                print(f"build_doc: dialog {title!r} is on screen; "
                                      f"closing it if still there in {grace:.0f} s",
                                      file=sys.stderr)
                            continue
                    keep.append((hwnd, title))
                found = keep
                for hwnd, title in found:
                    if title.startswith("Font Capture"):
                        what = "pressed OK on" if press_ok(hwnd) else "closed"
                        if what == "closed":
                            win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
                    else:
                        win32gui.PostMessage(hwnd, win32con.WM_CLOSE, 0, 0)
                        what = "cancelled"
                    self.handled.append(title)
                    print(f"build_doc: {what} the dialog {title!r}",
                          file=sys.stderr)
            except Exception:
                pass
            self._stop.wait(1.0)


def _export_via_pdfmaker(word, doc, pdf_path: Path):
    """Word's 'Save as Adobe PDF' (the Acrobat PDFMaker add-in), scripted —
    DOCUMENT_STANDARDS section 8's preferred export when Acrobat is healthy
    (named font subsets, bookmarks, links). Drives the add-in's IPDFMaker
    interface from the AdobePDFMakerForOffice type library: the add-in's
    Object exposes no type info, so pywin32 binds by interface IID.

    STATUS 27-Sep-2026: reachable and settable (prompt/view/progress off,
    output path honored), but CreatePDFEx launched the broken Acrobat.exe
    for font capture, that crashed into its modal box and Word never came
    back (RPC 'disconnected from its clients', no PDF). Opt-in only
    (DSES_PDF_ENGINE=pdfmaker) until Acrobat is repaired."""
    import time
    import pythoncom
    from win32com.client import gencache
    mod = gencache.EnsureModule('{EA8D486A-09E6-411F-B452-78F075ACC8CC}', 0, 1, 0)
    if mod is None:
        raise RuntimeError("AdobePDFMakerForOffice type library not registered")
    addin = word.COMAddIns.Item("PDFMaker.OfficeAddin")
    if not addin.Connect:
        raise RuntimeError("PDFMaker.OfficeAddin is not connected in Word")
    raw = addin.Object._oleobj_
    try:
        raw = raw.QueryInterface(mod.IPDFMaker.CLSID, pythoncom.IID_IDispatch)
    except pythoncom.com_error:
        pass
    maker = mod.IPDFMaker(raw)
    settings = maker.GetCurrentConversionSettings()
    if not isinstance(settings, mod.ISettings):
        settings = mod.ISettings(settings._oleobj_)
    if pdf_path.exists():
        pdf_path.unlink()
    settings.OutputPDFFileName = str(pdf_path.resolve())
    settings.PromptForPDFFilename = False
    settings.ViewPDFFile = False
    settings.ShouldShowProgressDialog = False
    maker.CreatePDFEx(settings, 0)
    for _ in range(120):
        if pdf_path.is_file() and pdf_path.stat().st_size > 0:
            return
        time.sleep(1.0)
    raise RuntimeError("PDFMaker produced no PDF")


def _convert_docx_to_pdf_word(docx_path: Path, pdf_path: Path):
    """Open the DOCX in Word, update every field (TOC + PAGE), save it, then
    write the PDF through the engine _pdf_engine() names:

      distiller (default when installed) Word -> PostScript -> acrodist.exe
                (_print_to_adobe_pdf): the daily output, named MinionPro/
                MyriadPro subsets. The attempt is bounded — a crashed
                font-capture helper is killed and the build falls back to
                msprint by itself.
      msprint   'Microsoft Print to PDF' via PrintOut/PrintToFile: silent, no
                Acrobat involvement; fonts embedded, text selectable
                (_print_to_ms_pdf).
      pdfmaker  Word's 'Save as Adobe PDF' add-in (_export_via_pdfmaker);
                falls back to msprint on failure.

    Word's own exporter (SaveAs/ExportAsFixedFormat FileFormat=17) is never
    used: it rasterizes the OpenType-PS house faces (tested 19-Aug-2026 and
    18-Sep-2026 — image-only pages, no text layer). A dialog sentinel runs
    during the print step so a modal print box cannot park the build on a
    human. Runs under build_lock (one Word/PDF build at a time on the
    machine, so another session's build waits instead of colliding), and
    under that lock kills orphaned invisible winword.exe processes and
    leftover Distillers first (a Word with a visible window is left alone)."""
    with build_lock(f"the PDF of {docx_path.name}"):
        return _convert_docx_to_pdf_word_locked(docx_path, pdf_path)


def _convert_docx_to_pdf_word_locked(docx_path: Path, pdf_path: Path):
    import time
    _cleanup_scratch()
    kill_stale_distillers("at build start")
    kill_acrotray("at build start")
    _kill_stale_invisible_word()
    _remove_stale_owner_file(docx_path)
    engine = _pdf_engine()
    t_start = time.time()
    word = _word_app()
    word_pid = _newest_winword_pid(t_start)
    word.Visible = False
    try:
        word.DisplayAlerts = 0          # wdAlertsNone
    except Exception:
        pass
    used = engine
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        try:
            # Update headers/footers + body fields. Run twice: the first
            # pass renders the TOC entries which can change page counts,
            # the second pass corrects the TOC's page numbers in light of
            # the new layout.
            for _ in range(2):
                doc.Fields.Update()
                # TablesOfContents may also need its own Update call on some
                # Word builds.
                try:
                    if doc.TablesOfContents.Count > 0:
                        doc.TablesOfContents(1).Update()
                except Exception:
                    pass
            # Word writes each TOC entry's paragraph mark in the THEME font
            # (asciiTheme=minorHAnsi, 12 pt) regardless of the toc styles, and
            # that font then rides into the PDF's embedded-font list (Cambria,
            # 2026-09-10). Pin the whole TOC range to the house heading face;
            # sizes and the level-1 bold still come from the toc styles.
            try:
                if doc.TablesOfContents.Count > 0:
                    doc.TablesOfContents(1).Range.Font.Name = FONT_HEAD
            except Exception:
                pass
            # Save the .docx so the populated TOC persists for future opens
            # in Word (otherwise the TOC reverts to the placeholder).
            doc.Save()
            with _DialogSentinel(word_pid):
                if engine == "distiller":
                    if not _adobe_pdf_printer_available():
                        raise RuntimeError("DSES_PDF_ENGINE=distiller but the "
                                           "'Adobe PDF' printer is not installed")
                    try:
                        _print_to_adobe_pdf(word, doc, docx_path, pdf_path)
                    except Exception as exc:
                        print("NOTE: the Distiller route failed "
                              f"({type(exc).__name__}: {exc}).\n"
                              "  Known causes: Acrobat.exe crashing in font "
                              "capture (the 'Font Capture: Windows - "
                              "Application Error' box), or the Adobe PDF "
                              "printer option\n  'Rely on system fonts only; "
                              "do not use document fonts' ticked again. "
                              "Falling back to 'Microsoft Print to PDF' "
                              "(fonts embedded, text selectable).",
                              file=sys.stderr)
                        _dismiss_adobe_pdf_dialog()
                        used = f"msprint (Distiller route failed: {exc})"
                        _print_to_ms_pdf(word, doc, pdf_path)
                elif engine == "pdfmaker":
                    try:
                        _export_via_pdfmaker(word, doc, pdf_path)
                    except Exception as exc:
                        print("WARNING: the PDFMaker route failed "
                              f"({type(exc).__name__}: {exc}). Falling back "
                              "to 'Microsoft Print to PDF'.", file=sys.stderr)
                        used = "msprint (PDFMaker route failed)"
                        _print_to_ms_pdf(word, doc, pdf_path)
                else:
                    _print_to_ms_pdf(word, doc, pdf_path)
        finally:
            doc.Close(SaveChanges=False)
    finally:
        word.Quit()
    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size // 1024} KB) "
          f"via {used}")


def main():
    global DOC_TITLE, DOC_SUBTITLE  # cover-page/header read these as globals
    import argparse
    ap = argparse.ArgumentParser(
        description="Build a DSES-styled DOCX + PDF from a Markdown source.")
    ap.add_argument('src', nargs='?', default=str(SRC),
                    help=f"Markdown source (default: {SRC})")
    ap.add_argument('--docx', default=None,
                    help="Keep the intermediate DOCX at this path (default: a "
                         "temp file, removed after the PDF is built).")
    ap.add_argument('--pdf', default=str(DST_PDF),
                    help=f"PDF output path (default: {DST_PDF})")
    ap.add_argument('--title', default=DOC_TITLE,
                    help=f"Cover-page title (default: {DOC_TITLE!r})")
    ap.add_argument('--subtitle', default=DOC_SUBTITLE,
                    help=f"Cover-page + header subtitle "
                         f"(default: {DOC_SUBTITLE!r})")
    ap.add_argument('--version', default=DOC_VERSION, dest='doc_version',
                    help="Cover-page version line. A value starting with 'v' "
                         "renders as 'Version X.Y.Z'; anything else (e.g. "
                         f"'Rev A') is printed verbatim. Default: {DOC_VERSION!r}")
    ap.add_argument('--header-logo', default=None,
                    help="PNG shown right-justified in the page header "
                         "(pages 2+; the cover is unaffected). Default: none.")
    ap.add_argument('--no-cover', action='store_true',
                    help="Skip the cover page (short memos / one-page "
                         "instruction sheets).")
    ap.add_argument('--no-toc', action='store_true',
                    help="Skip the table of contents.")
    ap.add_argument('--toc', choices=['field', 'static'], default=None,
                    dest='toc_mode',
                    help="TOC style: 'field' = real Word TOC field with page "
                         "numbers (evaluated by Word; the default on Windows), "
                         "'static' = pre-populated outline without page numbers "
                         "(renders anywhere; the default off Windows).")
    ap.add_argument('--force', action='store_true',
                    help="Overwrite the --docx target even if git reports it "
                         "modified (i.e., discard hand-made Word edits).")
    args = ap.parse_args()
    DOC_TITLE = args.title
    DOC_SUBTITLE = args.subtitle
    globals()['DOC_VERSION'] = args.doc_version
    globals()['HEADER_LOGO'] = args.header_logo

    src = Path(args.src)
    pdf_out = Path(args.pdf)
    if not src.exists():
        print(f"Source not found: {src}", file=sys.stderr)
        sys.exit(1)

    # The DOCX is a throwaway intermediate: write it to a temp file and remove
    # it after the PDF is built, unless --docx asked to keep it somewhere.
    import tempfile
    keep_docx = args.docx is not None
    if keep_docx:
        docx_out = Path(args.docx)
    else:
        docx_out = Path(tempfile.gettempdir()) / f"{pdf_out.stem}.docx"

    # GUARD: never clobber hand-made Word edits. Rebuilding regenerates the
    # DOCX from the .md, so any direct edits to a kept, git-tracked DOCX are
    # lost. If git reports the target modified vs HEAD, someone edited it
    # since the last build -- stop and make them port the edits into the .md
    # (or pass --force to overwrite deliberately). Born of the 2026-07-13
    # trip-report incident; do not remove.
    if keep_docx and docx_out.exists() and not args.force:
        import subprocess
        try:
            r = subprocess.run(
                ['git', 'status', '--porcelain', '--', str(docx_out)],
                capture_output=True, text=True, timeout=15)
            dirty = r.returncode == 0 and r.stdout.strip().startswith(' M')
        except Exception:
            dirty = False
        if dirty:
            print(f"REFUSING to overwrite {docx_out}:\n"
                  "  git says it was modified since the last commit — it "
                  "likely contains\n  hand-made Word edits that a rebuild "
                  "would destroy.\n"
                  "  Port those edits into the Markdown source first (or "
                  "commit the DOCX),\n  then rebuild. Pass --force to "
                  "overwrite anyway.", file=sys.stderr)
            sys.exit(3)

    md_to_docx(src, docx_out, cover=not args.no_cover, toc=not args.no_toc,
               toc_mode=args.toc_mode)
    try:
        convert_docx_to_pdf(docx_out, pdf_out)
    except Exception as exc:
        print(f"PDF conversion failed: {type(exc).__name__}: {exc}",
              file=sys.stderr)
        print(f"The intermediate DOCX is at {docx_out}; open it in "
              "LibreOffice/Word and export to PDF manually.", file=sys.stderr)
        sys.exit(2)
    finally:
        if not keep_docx:
            try:
                docx_out.unlink()
            except OSError:
                pass


if __name__ == '__main__':
    main()
