"""Tests for build_doc.py, the document builder.

    .conda\\python.exe test_build_doc.py

Born of two faults found 2026-10-04 in the published install guide:

* every guide printed the backticks of **`launcher.bat`**, because the inside
  of a bold span was emitted verbatim (emphasis must nest);
* at the 1.1.8 cut (3-Aug-2026) every code fence in Installing.md lost a
  backtick, and for two months each command block printed as run-on inline
  text with curly quotes and the language word glued on - including the
  desktop-shortcut command. Nothing in the build complained. The builder now
  refuses such a source, and the tests below pin how lists, code blocks and
  headings must come out.
"""
import sys
import tempfile
from pathlib import Path

import build_doc

F = "`" * 3                     # a code fence, kept out of this file's own text
B, I, C = {"bold": True}, {"italic": True}, {"code": True}

INLINE_CASES = [
    ("plain text", [("plain text", {})]),
    ("a **bold** word", [("a ", {}), ("bold", B), (" word", {})]),
    ("an *italic* word", [("an ", {}), ("italic", I), (" word", {})]),
    ("run `launcher.bat` now", [("run ", {}), ("launcher.bat", C), (" now", {})]),
    ("Double-click **`launcher.bat`** to start",
     [("Double-click ", {}), ("launcher.bat", {**B, **C}), (" to start", {})]),
    ("**run `a.ps1` or `b.bat` here**",
     [("run ", B), ("a.ps1", {**B, **C}), (" or ", B), ("b.bat", {**B, **C}), (" here", B)]),
    ("*see `settings.ini` first*",
     [("see ", I), ("settings.ini", {**I, **C}), (" first", I)]),
    ("**read [the guide](https://example.org/g.pdf)**",
     [("read ", B), ("the guide", B), (" (https://example.org/g.pdf)", {**B, **C})]),
    ("see [the guide](https://example.org/g.pdf).",
     [("see ", {}), ("the guide", {}), (" (https://example.org/g.pdf)", C), (".", {})]),
    ("[https://example.org](https://example.org)", [("https://example.org", C)]),
    ("type `**kwargs` and `*args`", [("type ", {}), ("**kwargs", C), (" and ", {}), ("*args", C)]),
    ("**A.** then `b` then *c*", [("A.", B), (" then ", {}), ("b", C), (" then ", {}), ("c", I)]),
]

DOC = f"""# Title

## 1. Section with `code` and **bold** in the heading

1. First step, download **`setup.exe`**.
   - **Primary:** <https://example.org/a>
   - **Backup:** <https://example.org/b>

   (An indented note under step one.)
2. Run it:

   {F}bash
   bash install.sh "my file"
     indented --flag
   {F}

   Accept the "license".
3. Third step.
4. Fourth step.

- A bullet with a block:
  {F}text
  copy "a" "b"
  {F}

A plain "quoted" paragraph.

{F}powershell
powershell -File .\\x.ps1
{F}

## 2. Another section

1. New list starts at one.
2. And two.
"""

checks = {"n": 0, "fail": 0}


def check(ok, label, detail=""):
    checks["n"] += 1
    if not ok:
        checks["fail"] += 1
    print("ok  " if ok else "FAIL", label, ("-> " + str(detail)) if (detail and not ok) else "")


def num_id(par):
    npr = par._p.pPr.numPr if par._p.pPr is not None else None
    return None if npr is None or npr.numId is None else npr.numId.val


def main():
    # ---- inline -----------------------------------------------------------
    for md, want in INLINE_CASES:
        got = [(t, s) for t, s in build_doc.parse_inline(md) if t]
        check(got == want, f"inline {md!r}", got)

    # ---- blocks -----------------------------------------------------------
    import docx
    tmp = Path(tempfile.mkdtemp(prefix="dses_build_doc_test_"))
    src, dst = tmp / "t.md", tmp / "t.docx"
    src.write_text(DOC, encoding="utf-8")
    build_doc.md_to_docx(src, dst, cover=False, toc=False)
    d = docx.Document(str(dst))

    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    check("1. Section with code and bold in the heading" in heads, "heading prints without markers", heads)

    nums = [(p.text, num_id(p)) for p in d.paragraphs if p.style.name == "List Number"]
    texts = [t for t, _ in nums]
    check(len(nums) == 6, "six numbered items", texts)
    first, second = nums[:4], nums[4:]
    check(all(n is not None for _, n in nums), "every numbered item carries a list id", nums)
    check(len({n for _, n in first}) == 1,
          "steps 1-4 are ONE list across the code block and the nested bullets", first)
    check(len({n for _, n in second}) == 1 and second[0][1] != first[0][1],
          "the list after the next heading restarts", second)

    blocks = [[p.text for p in t.cell(0, 0).paragraphs] for t in d.tables]
    check(blocks == [['bash install.sh "my file"', '  indented --flag'],
                     ['copy "a" "b"'],
                     ['powershell -File .\\x.ps1']],
          "three code blocks: own lines, indent removed, straight quotes, no language word", blocks)

    body = {p.text: p for p in d.paragraphs}
    lic = body.get("Accept the “license”.")
    note = body.get("(An indented note under step one.)")
    check(lic is not None and lic.paragraph_format.left_indent is not None
          and abs(lic.paragraph_format.left_indent.inches - 0.5) < 0.01,
          "paragraph under a list item keeps the item indent (and gets curly quotes)")
    check(note is not None and note.paragraph_format.left_indent is not None,
          "note under step one keeps the item indent")
    plain = body.get("A plain “quoted” paragraph.")
    check(plain is not None and plain.paragraph_format.left_indent is None,
          "an unindented paragraph is not indented")

    everything = [p.text for p in d.paragraphs] + [x for b in blocks for x in b]
    check(not any("`" in t for t in everything), "no backtick anywhere in the document",
          [t for t in everything if "`" in t])
    check(not any(t.startswith(("bash bash", "text copy", "powershell powershell")) for t in everything),
          "no language word glued onto a command")

    setup = [r for p in d.paragraphs for r in p.runs if r.text == "setup.exe"]
    check(len(setup) == 1 and setup[0].bold and setup[0].font.name == build_doc.FONT_MONO,
          "code inside bold is bold monospace")

    bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
    check("A bullet with a block:" in bullets, "a fence under a bullet is not swallowed into the bullet", bullets)

    # ---- the guard ----------------------------------------------------------
    bad = tmp / "bad.md"
    bad.write_text("# T\n\ntext\n\n``bash\nrm -rf x\n``\n", encoding="utf-8")
    try:
        build_doc.md_to_docx(bad, tmp / "bad.docx", cover=False, toc=False)
        check(False, "a two-backtick fence is refused")
    except SystemExit as exc:
        check("damaged code fence" in str(exc) and "5, 7" in str(exc),
              "a two-backtick fence is refused, with its line numbers", exc)
    ok_inline = tmp / "ok.md"
    ok_inline.write_text("# T\n\nuse ``double`` backticks inline, and `single`.\n", encoding="utf-8")
    try:
        build_doc.md_to_docx(ok_inline, tmp / "ok.docx", cover=False, toc=False)
        check(True, "inline double backticks do not trip the guard")
    except SystemExit as exc:
        check(False, "inline double backticks do not trip the guard", exc)

    print(f"{checks['n'] - checks['fail']} passed, {checks['fail']} failed")
    return 1 if checks["fail"] else 0


if __name__ == "__main__":
    sys.exit(main())
